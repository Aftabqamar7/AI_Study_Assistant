
import io
import re
import json
import html

import streamlit as st
from groq import Groq
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
from PIL import Image
import pytesseract

# ==================== CONFIGURATION ====================

MODEL = "openai/gpt-oss-120b"

st.set_page_config(
    page_title="Lumina AI Study Assistant",
    page_icon="✨",
    layout="wide",
)

# ==================== DESIGN ====================

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');

.stApp {
    background:
        radial-gradient(circle at top left, #252052, transparent 35%),
        radial-gradient(circle at top right, #123344, transparent 30%),
        #0a0d18;
    color: #f5f7ff;
    font-family: Inter, sans-serif;
}
[data-testid="stSidebar"] {
    background: #101529;
    border-right: 1px solid #29334d;
}
h1, h2, h3, h4 { color: #ffffff !important; }
p, label { color: #d7def3; }

.hero {
    padding: 30px;
    border-radius: 22px;
    border: 1px solid #494073;
    background: linear-gradient(120deg, #251d49, #112a3b);
    margin-bottom: 22px;
}
.hero-title {
    font-size: 2.5rem;
    font-weight: 800;
    color: white;
}
.hero-sub {
    color: #bbc7e4;
    font-size: 1rem;
}
.metric {
    background: #151c30;
    border: 1px solid #2c3855;
    padding: 18px;
    border-radius: 16px;
    margin-bottom: 12px;
}
.metric-label {
    color: #aab6d5;
    font-size: 0.85rem;
}
.metric-value {
    color: white;
    font-size: 1.5rem;
    font-weight: 800;
    margin-top: 8px;
    overflow-wrap: anywhere;
}
.panel {
    background: #141b2e;
    border: 1px solid #2b3652;
    border-radius: 17px;
    padding: 20px;
    margin-bottom: 14px;
}
.eyebrow {
    color: #b6a3ff;
    font-size: .78rem;
    font-weight: 800;
    letter-spacing: .12em;
    text-transform: uppercase;
}
div.stButton > button,
div.stFormSubmitButton > button {
    background: linear-gradient(100deg, #7955e8, #536fe5);
    color: white;
    border: 1px solid #6556b5;
    border-radius: 11px;
    font-weight: 700;
    min-height: 2.7rem;
}
div.stButton > button:hover,
div.stFormSubmitButton > button:hover {
    border-color: #22d3ee;
    color: white;
}
.stTextInput input, .stTextArea textarea,
[data-baseweb="select"] > div {
    background: #101629 !important;
    color: white !important;
    border-color: #303c58 !important;
    border-radius: 10px !important;
}
.stTabs [data-baseweb="tab-list"] {
    gap: 8px;
    background: #111629;
    border-radius: 12px;
    padding: 7px;
}
.stTabs [aria-selected="true"] {
    background: #302653 !important;
}
div[data-testid="stFileUploader"] {
    background: #12182b;
    border: 1px dashed #495578;
    border-radius: 14px;
    padding: 10px;
}
.small {
    color: #aab6d5;
    font-size: .88rem;
}
.question-card {
    background: #171f34;
    border: 1px solid #34415f;
    border-radius: 13px;
    padding: 16px;
    margin: 10px 0;
}
hr { border-color: #29344e; }
</style>
""", unsafe_allow_html=True)

# ==================== SESSION STATE ====================

defaults = {
    "document_text": "",
    "document_name": "",
    "chat_history": [],
    "quiz": [],
    "quiz_submitted": False,
    "quiz_answers": {},
    "quiz_score": None,
    "groq_api_key": "",
    "quiz_topic": "",
}

for key, value in defaults.items():
    if key not in st.session_state:
        st.session_state[key] = value

# ==================== DOCUMENT EXTRACTION ====================

def extract_document(uploaded_file):
    """Extract text from supported study materials."""
    name = uploaded_file.name.lower()
    content = uploaded_file.getvalue()

    # PDF
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(content))
        parts = []

        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                parts.append(f"Page {page_number}\n{text}")

        return "\n\n".join(parts).strip()

    # Word document
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(content))
        parts = []

        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text.strip())

        for table in doc.tables:
            for row in table.rows:
                parts.append(
                    " | ".join(cell.text for cell in row.cells)
                )

        return "\n".join(parts).strip()

    # PowerPoint
    if name.endswith(".pptx"):
        presentation = Presentation(io.BytesIO(content))
        slides_text = []

        for slide_number, slide in enumerate(
            presentation.slides, start=1
        ):
            slide_parts = []

            for shape in slide.shapes:
                if shape.has_text_frame and shape.text.strip():
                    slide_parts.append(shape.text.strip())

                if shape.has_table:
                    for row in shape.table.rows:
                        slide_parts.append(
                            " | ".join(cell.text for cell in row.cells)
                        )

            slides_text.append(
                f"Slide {slide_number}\n"
                + "\n".join(slide_parts)
            )

        return "\n\n".join(slides_text).strip()

    # Individual image / slide screenshot (OCR)
    if name.endswith((".png", ".jpg", ".jpeg", ".tif", ".tiff")):
        image = Image.open(io.BytesIO(content)).convert("RGB")
        return pytesseract.image_to_string(image).strip()

    # Text and Markdown
    if name.endswith((".txt", ".md")):
        return content.decode("utf-8", errors="replace").strip()

    raise ValueError(
        "Unsupported file. Upload PDF, DOCX, PPTX, TXT, MD, "
        "PNG, JPG, JPEG, TIF, or TIFF."
    )


# ==================== DOCUMENT RETRIEVAL ====================

def retrieve_context(question, max_chars=12000):
    text = st.session_state.document_text

    if not text:
        return ""

    paragraphs = [
        part.strip()
        for part in re.split(r"\n+", text)
        if part.strip()
    ]

    keywords = {
        word.lower()
        for word in re.findall(r"\w+", question)
        if len(word) > 2
    }

    scored = []

    for index, paragraph in enumerate(paragraphs):
        words = {
            word.lower()
            for word in re.findall(r"\w+", paragraph)
            if len(word) > 2
        }

        score = len(keywords.intersection(words))

        if score:
            scored.append((score, index, paragraph))

    scored.sort(key=lambda item: (-item[0], item[1]))

    if scored:
        selected = [item[2] for item in scored[:20]]
    else:
        selected = paragraphs[:10]

    return "\n\n".join(selected)[:max_chars]


# ==================== GROQ AI ====================

def get_ai_response(messages, temperature=0.3):
    api_key = st.session_state.get("groq_api_key", "").strip()

    if not api_key:
        raise ValueError(
            "Enter your Groq API key in the sidebar."
        )

    client = Groq(api_key=api_key)

    response = client.chat.completions.create(
        model=MODEL,
        messages=messages,
        temperature=temperature,
    )

    return response.choices[0].message.content or ""


# ==================== STUDY CHAT ====================

def explain_topic(question, language, style):
    context = retrieve_context(question)

    languages = {
        "English": "Answer in simple, clear English.",
        "Urdu": "جواب آسان اور واضح اردو میں دیں۔ اردو رسم الخط استعمال کریں۔",
        "Roman Urdu": "Answer in easy Roman Urdu.",
        "English + Urdu": (
            "First answer in simple English, then explain in Urdu script."
        ),
    }

    styles = {
        "Simple Explanation": (
            "Explain for a beginner and include a simple example."
        ),
        "Exam Notes": (
            "Use headings, short bullet points, and key exam points."
        ),
        "Detailed Explanation": (
            "Explain clearly with useful details and examples."
        ),
    }

    messages = [
        {
            "role": "system",
            "content": (
                "You are Lumina, a helpful study tutor. "
                "Use the uploaded document as your main source. "
                "Do not invent facts or claim unsupported information "
                "comes from the document. If the document lacks the "
                "answer, say so clearly. "
                + languages[language]
                + " "
                + styles[style]
            ),
        },
        {
            "role": "user",
            "content": (
                f"DOCUMENT CONTEXT:\n{context}\n\n"
                f"STUDENT QUESTION:\n{question}"
            ),
        },
    ]

    return get_ai_response(messages)


# ==================== QUIZ GENERATION ====================

def generate_quiz(topic, count, difficulty, language):
    context = retrieve_context(
        topic or "main concepts",
        max_chars=14000,
    )

    prompt = f"""
Create exactly {count} multiple-choice questions based on the
document context below.

Difficulty: {difficulty}
Language: {language}

Return ONLY a valid JSON array with this structure:

[
  {{
    "question": "Question text",
    "options": ["Option A", "Option B", "Option C", "Option D"],
    "answer": 0,
    "explanation": "Explain why the answer is correct."
  }}
]

The answer must be an integer from 0 to 3.
Use clear questions based on the supplied material.
Do not use Markdown code fences.
For Urdu, use Urdu script. For Roman Urdu, use Roman Urdu.

DOCUMENT CONTEXT:
{context}

TOPIC:
{topic or "All main concepts"}
"""

    raw = get_ai_response(
        [
            {
                "role": "system",
                "content": (
                    "You create educational multiple-choice quizzes. "
                    "Return only valid JSON."
                ),
            },
            {"role": "user", "content": prompt},
        ],
        temperature=0.3,
    )

    raw = re.sub(
        r"^```(?:json)?\s*|\s*```$",
        "",
        raw.strip(),
        flags=re.IGNORECASE,
    )

    match = re.search(r"\[.*\]", raw, flags=re.DOTALL)

    if not match:
        raise ValueError(
            "The AI did not return valid quiz data. Please try again."
        )

    questions = json.loads(match.group(0))
    valid_questions = []

    for item in questions:
        options = item.get("options", [])
        answer = item.get("answer")

        if (
            isinstance(item.get("question"), str)
            and isinstance(options, list)
            and len(options) == 4
            and isinstance(answer, int)
            and not isinstance(answer, bool)
            and 0 <= answer <= 3
        ):
            valid_questions.append({
                "question": item["question"],
                "options": [str(option) for option in options],
                "answer": answer,
                "explanation": str(item.get("explanation", "")),
            })

    if not valid_questions:
        raise ValueError("No valid quiz questions were generated.")

    return valid_questions[:count]


# ==================== SIDEBAR ====================

with st.sidebar:
    st.markdown("## ✨ Lumina")
    st.caption("Your personal AI study space")
    st.divider()

    st.markdown("### 🔐 Connect Groq")

    st.text_input(
        "Your Groq API key",
        type="password",
        key="groq_api_key",
        placeholder="Enter your Groq API key",
        help="Get your key from console.groq.com.",
    )

    st.caption("Keep your API key private. Never commit it to GitHub.")

    st.divider()
    st.markdown("### 📚 Upload study material")

    uploaded_file = st.file_uploader(
        "Choose your study material",
        type=[
            "pdf", "docx", "pptx", "txt", "md",
            "png", "jpg", "jpeg", "tif", "tiff"
        ],
        help=(
            "Upload a PDF, PowerPoint, Word document, text file, "
            "Markdown file, or a single image."
        ),
    )

    if st.button(
        "📥 Load Document",
        use_container_width=True,
    ):
        if uploaded_file is None:
            st.warning("Please choose a file first.")
        else:
            try:
                with st.spinner("Reading your study material..."):
                    extracted = extract_document(uploaded_file)

                if not extracted.strip():
                    st.error(
                        "No readable text was found. "
                        "Scanned PDFs or images may require OCR."
                    )
                else:
                    st.session_state.document_text = extracted
                    st.session_state.document_name = uploaded_file.name
                    st.session_state.chat_history = []
                    st.session_state.quiz = []
                    st.session_state.quiz_answers = {}
                    st.session_state.quiz_submitted = False
                    st.session_state.quiz_score = None

                    st.success("Study material loaded successfully!")

            except Exception as error:
                st.error(f"Could not read this file: {error}")

    if st.session_state.document_text:
        st.markdown("**Current document**")
        st.write(st.session_state.document_name)
        st.caption(
            f"{len(st.session_state.document_text):,} characters extracted"
        )

        with st.expander("Preview extracted text"):
            st.text(
                st.session_state.document_text[:2500]
            )

        if st.button(
            "🗑️ Remove Document",
            use_container_width=True,
        ):
            st.session_state.document_text = ""
            st.session_state.document_name = ""
            st.session_state.chat_history = []
            st.session_state.quiz = []
            st.session_state.quiz_answers = {}
            st.session_state.quiz_submitted = False
            st.session_state.quiz_score = None
            st.rerun()

    st.divider()
    st.markdown("### ⚙️ Active model")
    st.code(MODEL)
    st.caption("Fixed model. No Llama model options.")


# ==================== MAIN HEADER ====================

st.markdown("""
<div class="hero">
    <div class="eyebrow">YOUR KNOWLEDGE, MADE CLEARER</div>
    <div class="hero-title">Study smarter with Lumina ✨</div>
    <p class="hero-sub">
        Understand your notes, ask questions, prepare for exams,
        and challenge yourself with interactive quizzes.
    </p>
    <span>
        📚 Document AI　·　🧠 Quiz Arena　·　🌐 Multilingual Learning
    </span>
</div>
""", unsafe_allow_html=True)

# ==================== DASHBOARD ====================

col1, col2, col3 = st.columns(3)

with col1:
    st.markdown(
        f"""
        <div class="metric">
            <div class="metric-label">📄 STUDY MATERIAL</div>
            <div class="metric-value">
                {"Ready" if st.session_state.document_text else "Not Added"}
            </div>
            <div class="small">
                {html.escape(st.session_state.document_name or "Upload from sidebar")}
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

with col2:
    st.markdown(
        f"""
        <div class="metric">
            <div class="metric-label">💬 QUESTIONS ASKED</div>
            <div class="metric-value">
                {len(st.session_state.chat_history)}
            </div>
            <div class="small">Your study conversations</div>
        </div>
        """,
        unsafe_allow_html=True,
    )

with col3:
    score = st.session_state.quiz_score
    score_text = f"{score}%" if score is not None else "—"

    st.markdown(
        f"""
        <div class="metric">
            <div class="metric-label">🏆 LATEST QUIZ SCORE</div>
            <div class="metric-value">{score_text}</div>
            <div class="small">Practice and improve</div>
        </div>
        """,
        unsafe_allow_html=True,
    )


# ==================== APPLICATION TABS ====================

study_tab, quiz_tab, guide_tab = st.tabs([
    "💬 Study Chat",
    "🏆 Quiz Arena",
    "🧭 Quick Start",
])


# ==================== STUDY CHAT ====================

with study_tab:
    st.markdown("## 💬 Ask your study material")
    st.caption(
        "Ask questions about your uploaded notes and receive "
        "student-friendly explanations."
    )

    left, right = st.columns(2)

    with left:
        answer_language = st.selectbox(
            "Answer language",
            ["English", "Urdu", "Roman Urdu", "English + Urdu"],
        )

    with right:
        explanation_style = st.selectbox(
            "Explanation style",
            [
                "Simple Explanation",
                "Exam Notes",
                "Detailed Explanation",
            ],
        )

    with st.form("study_question_form"):
        user_question = st.text_area(
            "Your question",
            placeholder=(
                "Example: Explain the main concept in simple words..."
            ),
            height=100,
        )

        ask_button = st.form_submit_button(
            "✨ Explain This Topic",
            use_container_width=True,
        )

    if ask_button:
        if not st.session_state.groq_api_key.strip():
            st.warning("Enter your Groq API key in the sidebar.")
        elif not st.session_state.document_text:
            st.warning("Upload and load a document first.")
        elif not user_question.strip():
            st.warning("Please type a question.")
        else:
            with st.spinner("Lumina is preparing your explanation..."):
                try:
                    answer = explain_topic(
                        user_question.strip(),
                        answer_language,
                        explanation_style,
                    )

                    st.session_state.chat_history.append({
                        "question": user_question.strip(),
                        "answer": answer,
                    })

                except Exception as error:
                    st.error(f"Could not generate answer: {error}")

    if st.session_state.chat_history:
        st.divider()
        st.markdown("### 📖 Your study conversation")

        for item in reversed(st.session_state.chat_history):
            with st.container():
                st.markdown("**🙋 Your question**")
                st.write(item["question"])
                st.markdown("**✨ Lumina explains**")
                st.markdown(item["answer"])
                st.divider()

        if st.button("🧹 Clear Conversation"):
            st.session_state.chat_history = []
            st.rerun()

    elif not st.session_state.document_text:
        st.info(
            "👈 Enter your Groq API key and load study material "
            "using the sidebar to begin."
        )


# ==================== QUIZ ARENA ====================

with quiz_tab:
    st.markdown("## 🏆 Quiz Arena")
    st.write(
        "Turn your study material into a challenge. Answer the questions, "
        "submit your quiz, and review your mistakes."
    )

    st.markdown("""
    <div class="panel">
        <div class="eyebrow">BUILD YOUR CHALLENGE</div>
        <h3>🎯 Ready to test your knowledge?</h3>
        Choose your topic, difficulty, language, and number of questions.
    </div>
    """, unsafe_allow_html=True)

    q1, q2, q3 = st.columns([1.5, 1, 1])

    with q1:
        quiz_topic = st.text_input(
            "Topic (optional)",
            placeholder="e.g. AES encryption",
            key="quiz_topic",
        )

    with q2:
        question_count = st.selectbox(
            "Number of questions",
            [5, 10, 15],
        )

    with q3:
        difficulty = st.selectbox(
            "Difficulty",
            ["Easy", "Medium", "Hard"],
            index=1,
        )

    quiz_language = st.selectbox(
        "Quiz language",
        ["English", "Urdu", "Roman Urdu"],
    )

    if st.button(
        "⚡ Generate My Quiz",
        use_container_width=True,
    ):
        if not st.session_state.groq_api_key.strip():
            st.warning("Enter your Groq API key in the sidebar.")
        elif not st.session_state.document_text:
            st.warning("Load a document before generating a quiz.")
        else:
            with st.spinner("Creating your quiz..."):
                try:
                    questions = generate_quiz(
                        quiz_topic,
                        question_count,
                        difficulty,
                        quiz_language,
                    )

                    st.session_state.quiz = questions
                    st.session_state.quiz_answers = {}
                    st.session_state.quiz_submitted = False
                    st.session_state.quiz_score = None

                    st.success(
                        f"{len(questions)} questions are ready!"
                    )

                except Exception as error:
                    st.error(f"Could not generate quiz: {error}")

    if st.session_state.quiz:
        st.divider()
        st.markdown(
            f"### 🎯 {quiz_topic or 'Document Knowledge'}"
        )

        total_questions = len(st.session_state.quiz)
        answered_count = len(st.session_state.quiz_answers)

        if not st.session_state.quiz_submitted:
            st.progress(
                min(answered_count / total_questions, 1.0),
                text=f"{answered_count} of {total_questions} answered",
            )

        with st.form("quiz_form"):
            selected_answers = {}

            for index, item in enumerate(st.session_state.quiz):
                st.markdown(
                    f"""
                    <div class="question-card">
                        <div class="eyebrow">
                            QUESTION {index + 1} OF {total_questions}
                        </div>
                        <h4>{html.escape(item["question"])}</h4>
                    </div>
                    """,
                    unsafe_allow_html=True,
                )

                previous_answer = st.session_state.quiz_answers.get(
                    str(index)
                )

                selected = st.radio(
                    f"Answer {index + 1}",
                    options=[0, 1, 2, 3],
                    format_func=lambda choice, options=item["options"]: (
                        options[choice]
                    ),
                    index=previous_answer,
                    key=f"quiz_answer_{index}",
                    disabled=st.session_state.quiz_submitted,
                    label_visibility="collapsed",
                )

                selected_answers[str(index)] = selected

            submit_answers = False

            if not st.session_state.quiz_submitted:
                submit_answers = st.form_submit_button(
                    "🏁 Submit Answers",
                    use_container_width=True,
                )

        if submit_answers:
            st.session_state.quiz_answers = selected_answers

            correct_count = sum(
                selected_answers.get(str(i)) == item["answer"]
                for i, item in enumerate(st.session_state.quiz)
            )

            st.session_state.quiz_score = round(
                100 * correct_count / total_questions
            )
            st.session_state.quiz_submitted = True
            st.rerun()

        if st.session_state.quiz_submitted:
            correct_count = sum(
                st.session_state.quiz_answers.get(str(i)) == item["answer"]
                for i, item in enumerate(st.session_state.quiz)
            )

            percentage = round(
                100 * correct_count / total_questions
            )

            st.divider()
            st.markdown("## 🎉 Your Results")

            result1, result2, result3 = st.columns(3)

            result1.metric(
                "Correct Answers",
                f"{correct_count}/{total_questions}",
            )
            result2.metric("Accuracy", f"{percentage}%")
            result3.metric(
                "Performance",
                (
                    "Excellent!"
                    if percentage >= 80
                    else "Good effort"
                    if percentage >= 50
                    else "Keep practicing"
                ),
            )

            st.progress(percentage / 100)

            st.markdown("### 📖 Review Your Answers")

            for index, item in enumerate(st.session_state.quiz):
                chosen = st.session_state.quiz_answers.get(str(index))
                correct = chosen == item["answer"]

                with st.expander(
                    f'{"✅" if correct else "❌"} Question {index + 1}'
                ):
                    st.write(item["question"])

                    if chosen is not None:
                        st.write(
                            f"**Your answer:** {item['options'][chosen]}"
                        )
                    else:
                        st.write("**Your answer:** Not answered")

                    st.write(
                        f"**Correct answer:** "
                        f"{item['options'][item['answer']]}"
                    )
                    st.write(
                        f"**Explanation:** {item['explanation']}"
                    )

            if st.button(
                "🔁 Generate Another Quiz",
                use_container_width=True,
            ):
                st.session_state.quiz = []
                st.session_state.quiz_answers = {}
                st.session_state.quiz_submitted = False
                st.session_state.quiz_score = None
                st.rerun()

    else:
        st.markdown("""
        <div class="panel">
            <h3>✨ Your next challenge awaits</h3>
            <p>
                1. Load your study material.<br>
                2. Choose your quiz settings.<br>
                3. Click Generate My Quiz.<br>
                4. Submit your answers to see your score.
            </p>
        </div>
        """, unsafe_allow_html=True)


# ==================== QUICK START ====================

with guide_tab:
    st.markdown("## 🧭 Get Started")

    a, b, c = st.columns(3)

    with a:
        st.markdown("""
        <div class="panel">
            <div class="eyebrow">STEP 01</div>
            <h3>🔐 Connect</h3>
            Enter your Groq API key in the sidebar.
        </div>
        """, unsafe_allow_html=True)

    with b:
        st.markdown("""
        <div class="panel">
            <div class="eyebrow">STEP 02</div>
            <h3>📚 Upload</h3>
            Upload PDF, PowerPoint, Word, TXT, Markdown, or images.
        </div>
        """, unsafe_allow_html=True)

    with c:
        st.markdown("""
        <div class="panel">
            <div class="eyebrow">STEP 03</div>
            <h3>🚀 Practice</h3>
            Ask questions or test yourself in Quiz Arena.
        </div>
        """, unsafe_allow_html=True)

    st.markdown("### Supported file types")

    st.write(
        "PDF · PPTX · DOCX · TXT · MD · PNG · JPG · JPEG · TIF · TIFF"
    )

    st.markdown("### Active model")
    st.code(MODEL)

    st.caption(
        "Text is extracted from uploaded materials. OCR works on images, "
        "but accuracy depends on image quality. Scanned PDFs may need OCR."
    )


# ==================== FOOTER ====================

st.divider()

st.markdown(
    """
    <div style="text-align:center;color:#8490ae;padding:12px;">
        ✨ Lumina AI Study Assistant · Learn with clarity.
    </div>
    """,
    unsafe_allow_html=True,
)

